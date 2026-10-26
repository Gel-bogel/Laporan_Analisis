"""Pembacaan file Excel (.xlsx/.xls) menjadi matriks sel mentah, lintas OS.

- .xlsx  → openpyxl (read_only, data_only)
- .xls   → xlrd; bila gagal → konversi LibreOffice (soffice --headless --convert-to xlsx)
File sementara memakai tempfile.gettempdir() (aman untuk Windows).
"""
from __future__ import annotations

import hashlib
import io
import os
import shutil
import subprocess
import tempfile
import uuid
from typing import Dict, List, Optional

import numpy as np
import pandas as pd

import config

ERROR_MARK = "#ERR"


def file_hash(data: bytes) -> str:
    return hashlib.md5(data).hexdigest()


def ext_of(filename: str) -> str:
    return os.path.splitext(filename or "")[1].lower()


# --------------------------------------------------------------------------- LibreOffice
def _find_soffice() -> Optional[str]:
    for c in config.SOFFICE_CANDIDATES:
        p = shutil.which(c) or (c if os.path.isfile(c) else None)
        if p:
            return p
    return None


def soffice_convert(data: bytes, filename: str, target: str = "xlsx") -> bytes:
    exe = _find_soffice()
    if not exe:
        raise RuntimeError("LibreOffice (soffice) tidak ditemukan untuk konversi file .xls.")
    work = os.path.join(tempfile.gettempdir(), f"pln_conv_{uuid.uuid4().hex}")
    os.makedirs(work, exist_ok=True)
    try:
        src = os.path.join(work, "input" + (ext_of(filename) or ".xls"))
        with open(src, "wb") as fh:
            fh.write(data)
        profile = "file:///" + os.path.join(work, "profile").replace("\\", "/").lstrip("/")
        cmd = [exe, f"-env:UserInstallation={profile}", "--headless",
               "--convert-to", target, "--outdir", work, src]
        subprocess.run(cmd, check=True, capture_output=True, timeout=config.SOFFICE_TIMEOUT_SEC)
        out = os.path.join(work, "input." + target)
        with open(out, "rb") as fh:
            return fh.read()
    finally:
        shutil.rmtree(work, ignore_errors=True)


# --------------------------------------------------------------------------- pembaca sheet
def _xlsx_sheet_names(data: bytes) -> List[str]:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _xlsx_matrices(data: bytes, sheets: Optional[List[str]], max_cols: Optional[int]) -> Dict[str, pd.DataFrame]:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = {}
    try:
        for name in wb.sheetnames:
            if sheets is not None and name not in sheets:
                continue
            ws = wb[name]
            rows = []
            for r in ws.iter_rows(values_only=True, max_col=max_cols):
                rows.append(list(r))
            width = max((len(r) for r in rows), default=0)
            rows = [r + [None] * (width - len(r)) for r in rows]
            out[name] = pd.DataFrame(rows, dtype=object)
    finally:
        wb.close()
    return out


def _xls_matrices(data: bytes, sheets: Optional[List[str]], max_cols: Optional[int]) -> Dict[str, pd.DataFrame]:
    import xlrd
    wb = xlrd.open_workbook(file_contents=data)
    out = {}
    for sh in wb.sheets():
        if sheets is not None and sh.name not in sheets:
            continue
        ncols = sh.ncols if max_cols is None else min(sh.ncols, max_cols)
        rows = []
        for r in range(sh.nrows):
            row = []
            for c in range(ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_ERROR:
                    row.append(ERROR_MARK)
                elif cell.ctype == xlrd.XL_CELL_EMPTY or cell.ctype == xlrd.XL_CELL_BLANK:
                    row.append(None)
                else:
                    row.append(cell.value)
            rows.append(row)
        out[sh.name] = pd.DataFrame(rows, dtype=object)
    return out


def read_matrices(data: bytes, filename: str, sheets: Optional[List[str]] = None,
                  max_cols: Optional[int] = None) -> Dict[str, pd.DataFrame]:
    """Semua (atau sebagian) sheet → {nama_sheet: DataFrame sel mentah (header=None)}."""
    ext = ext_of(filename)
    if ext == ".xls":
        try:
            return _xls_matrices(data, sheets, max_cols)
        except Exception:
            converted = soffice_convert(data, filename, "xlsx")
            return _xlsx_matrices(converted, sheets, max_cols)
    return _xlsx_matrices(data, sheets, max_cols)


def sheet_names(data: bytes, filename: str) -> List[str]:
    if ext_of(filename) == ".xls":
        try:
            import xlrd
            return xlrd.open_workbook(file_contents=data, on_demand=True).sheet_names()
        except Exception:
            return _xlsx_sheet_names(soffice_convert(data, filename, "xlsx"))
    return _xlsx_sheet_names(data)


def matrix_to_table(mat: pd.DataFrame, header_row: int) -> pd.DataFrame:
    """Matriks mentah → DataFrame dengan header di baris header_row (nama kolom unik)."""
    raw_cols = list(mat.iloc[header_row])
    cols, seen = [], {}
    for i, c in enumerate(raw_cols):
        name = str(c).strip() if c is not None and not (isinstance(c, float) and np.isnan(c)) else f"_kol{i + 1}"
        if name in seen:
            seen[name] += 1
            name = f"{name}.{seen[name]}"
        else:
            seen[name] = 0
        cols.append(name)
    body = mat.iloc[header_row + 1:].copy()
    body.columns = cols
    body = body.dropna(how="all")
    return body.reset_index(drop=True)


def cell_text(v) -> str:
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)
